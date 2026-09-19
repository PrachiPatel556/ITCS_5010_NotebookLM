"""Check that supported document parsers extract actual slide and page text."""

from __future__ import annotations

import importlib.util
import unittest
from io import BytesIO

from notebooklm.ingest import _extract_pdf, _extract_pptx, ingest_url
from notebooklm.storage import NotebookStore
from pathlib import Path
from unittest.mock import patch
from uuid import uuid4


def _one_page_pdf(text: str) -> bytes:
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        f"<< /Length {len(stream)} >>\nstream\n".encode("ascii") + stream + b"\nendstream",
    ]
    data = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for number, body in enumerate(objects, 1):
        offsets.append(len(data))
        data.extend(f"{number} 0 obj\n".encode("ascii") + body + b"\nendobj\n")
    start_xref = len(data)
    data.extend(f"xref\n0 {len(offsets)}\n0000000000 65535 f \n".encode("ascii"))
    for offset in offsets[1:]:
        data.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    data.extend(
        f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{start_xref}\n%%EOF\n".encode("ascii")
    )
    return bytes(data)


class ExtractorTests(unittest.TestCase):
    @unittest.skipUnless(importlib.util.find_spec("pypdf"), "pypdf is not installed")
    def test_pdf_text_is_extracted(self) -> None:
        result = _extract_pdf(_one_page_pdf("The deadline is Friday"))
        self.assertIn("[Page 1]", result)
        self.assertIn("The deadline is Friday", result)

    @unittest.skipUnless(importlib.util.find_spec("pptx"), "python-pptx is not installed")
    def test_pptx_text_is_extracted(self) -> None:
        from pptx import Presentation
        from pptx.util import Inches

        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text = "Study topic: climate"
        output = BytesIO()
        presentation.save(output)
        result = _extract_pptx(output.getvalue())
        self.assertIn("[Slide 1]", result)
        self.assertIn("Study topic: climate", result)

    @unittest.skipUnless(importlib.util.find_spec("bs4"), "beautifulsoup4 is not installed")
    def test_url_html_excludes_scripts_and_is_indexed(self) -> None:
        root = Path(__file__).resolve().parent / ".tmp"
        root.mkdir(exist_ok=True)
        db_path = root / f"url-{uuid4().hex}.db"
        self.addCleanup(db_path.unlink, missing_ok=True)
        store = NotebookStore(db_path)
        notebook = store.create_notebook("Web notes")
        body = b"<html><head><title>Useful page</title></head><body><p>The target is 35 percent.</p><script>ignore me</script></body></html>"
        with patch("notebooklm.ingest._download_url", return_value=("https://example.com/", body, "text/html")), patch(
            "notebooklm.ingest.embed_texts", return_value=[[1.0, 0.0]]
        ):
            source = ingest_url(store, notebook["id"], "https://example.com/")
        self.assertEqual(source["name"], "Useful page")
        self.assertIn("35 percent", store.list_chunks(notebook["id"])[0]["text"])
        self.assertNotIn("ignore me", store.list_chunks(notebook["id"])[0]["text"])


if __name__ == "__main__":
    unittest.main()
