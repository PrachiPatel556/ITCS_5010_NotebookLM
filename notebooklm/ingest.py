"""Extract, chunk, embed, and store PDF, PowerPoint, text, and web sources."""

from __future__ import annotations

import ipaddress
import http.client
import re
import socket
import time
from io import BytesIO
from pathlib import Path
from typing import Any, Callable, Sequence
from urllib.parse import urljoin, urlparse

from .retrieval import DEFAULT_MODEL, embed_texts
from .storage import NotebookStore


MAX_FILE_BYTES = 30 * 1024 * 1024
MAX_WEB_BYTES = 5 * 1024 * 1024
MAX_EXTRACTED_CHARS = 500_000
MAX_CHUNKS = 400


def _normalize_text(text: str) -> str:
    lines = [re.sub(r"[ \t]+", " ", line).strip()
             for line in text.replace("\r\n", "\n").replace("\r", "\n").split("\n")]
    return re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip()


def _decode_text(data: bytes) -> str:
    """Decode common document/web encodings without optional detector packages."""
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("cp1252", errors="replace")


def chunk_text(text: str, chunk_size: int = 1200, overlap: int = 200) -> list[str]:
    """Split on words into approximately character-sized overlapping chunks."""
    if chunk_size < 100 or overlap < 0 or overlap >= chunk_size:
        raise ValueError("Use chunk_size >= 100 and 0 <= overlap < chunk_size")
    words = _normalize_text(text).split()
    chunks: list[str] = []
    start = 0
    while start < len(words):
        end = start
        length = 0
        while end < len(words) and (length + len(words[end]) + (end > start) <= chunk_size or end == start):
            length += len(words[end]) + (end > start)
            end += 1
        chunks.append(" ".join(words[start:end]))
        if end == len(words):
            break
        overlap_start = end
        overlap_length = 0
        while overlap_start > start + 1 and overlap_length < overlap:
            overlap_start -= 1
            overlap_length += len(words[overlap_start]) + 1
        start = max(start + 1, overlap_start)
    return chunks


def _extract_pdf(data: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise RuntimeError("Install pypdf to ingest PDF files") from exc
    try:
        reader = PdfReader(BytesIO(data))
        pages = [f"[Page {number}]\n{extracted}"
                 for number, page in enumerate(reader.pages, start=1)
                 if (extracted := (page.extract_text() or "").strip())]
    except Exception as exc:
        raise ValueError(f"Could not read PDF: {exc}") from exc
    return _normalize_text("\n\n".join(pages))


def _extract_pptx(data: bytes) -> str:
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise RuntimeError("Install python-pptx to ingest PowerPoint files") from exc
    try:
        presentation = Presentation(BytesIO(data))
        slides = []
        for number, slide in enumerate(presentation.slides, start=1):
            parts = []
            for shape in slide.shapes:
                if shape.has_text_frame and shape.text.strip():
                    parts.append(shape.text)
                if shape.has_table:
                    parts.extend(" | ".join(cell.text for cell in row.cells)
                                 for row in shape.table.rows)
            if any(part.strip() for part in parts):
                slides.append(f"[Slide {number}]\n" + "\n".join(parts))
    except Exception as exc:
        raise ValueError(f"Could not read PowerPoint file: {exc}") from exc
    return _normalize_text("\n\n".join(slides))


def _extract_file(path: Path) -> tuple[str, str]:
    if not path.is_file():
        raise FileNotFoundError(f"Source file does not exist: {path}")
    if path.stat().st_size > MAX_FILE_BYTES:
        raise ValueError("Source file exceeds the 30 MB upload limit")
    extension = path.suffix.lower()
    if extension not in {".pdf", ".pptx", ".txt"}:
        raise ValueError("Supported file types are PDF, PPTX, and TXT")
    data = path.read_bytes()
    if extension == ".pdf":
        return "pdf", _extract_pdf(data)
    if extension == ".pptx":
        return "pptx", _extract_pptx(data)
    return "txt", _normalize_text(_decode_text(data))


def ingest_text(store: NotebookStore, notebook_id: str, name: str, text: str,
                kind: str = "txt", uri: str | None = None,
                model_name: str = DEFAULT_MODEL,
                embedder: Callable[[Sequence[str]], Sequence[Sequence[float]]] | None = None
                ) -> dict[str, Any]:
    """Index extracted text. An embedder can be injected for tests."""
    normalized = _normalize_text(text)
    if not normalized:
        raise ValueError("Source contains no extractable text (scanned PDFs require OCR)")
    if len(normalized) > MAX_EXTRACTED_CHARS:
        raise ValueError("Extracted source exceeds the 500,000 character limit")
    chunks = chunk_text(normalized)
    if len(chunks) > MAX_CHUNKS:
        raise ValueError("Source exceeds the 400 chunk limit")
    vectors = embedder(chunks) if embedder is not None else embed_texts(chunks, model_name)
    return store.add_source(notebook_id, name, kind, normalized, chunks, vectors, uri)


def ingest_file(store: NotebookStore, notebook_id: str, file_path: str | Path,
                filename: str | None = None, model_name: str = DEFAULT_MODEL) -> dict[str, Any]:
    path = Path(file_path)
    kind, text = _extract_file(path)
    return ingest_text(store, notebook_id, filename or path.name, text, kind, model_name=model_name)


def _check_public_http_url(url: str) -> tuple[str, int, str]:
    """Resolve a public host and return the exact IP to use for the connection."""
    parsed = urlparse(url)
    if parsed.scheme not in {"http", "https"} or not parsed.hostname:
        raise ValueError("Enter a valid public HTTP or HTTPS URL")
    if parsed.username or parsed.password:
        raise ValueError("URLs containing credentials are not supported")
    try:
        port = parsed.port or (443 if parsed.scheme == "https" else 80)
    except ValueError as exc:
        raise ValueError("URL has an invalid port") from exc
    host = parsed.hostname.encode("idna").decode("ascii")
    try:
        addresses = socket.getaddrinfo(host, port, type=socket.SOCK_STREAM)
    except socket.gaierror as exc:
        raise ValueError(f"Could not resolve URL host: {host}") from exc
    if not addresses:
        raise ValueError(f"Could not resolve URL host: {host}")
    for address in addresses:
        ip = ipaddress.ip_address(address[4][0])
        if not ip.is_global:
            raise ValueError("URLs on private or local networks are not supported")
    return host, port, str(ipaddress.ip_address(addresses[0][4][0]))


class _PinnedHTTPConnection(http.client.HTTPConnection):
    def __init__(self, host: str, port: int, ip: str) -> None:
        super().__init__(host, port, timeout=5)
        self._pinned_ip = ip

    def connect(self) -> None:
        self.sock = socket.create_connection(
            (self._pinned_ip, self.port), self.timeout, self.source_address
        )


class _PinnedHTTPSConnection(http.client.HTTPSConnection):
    def __init__(self, host: str, port: int, ip: str) -> None:
        super().__init__(host, port, timeout=5)
        self._pinned_ip = ip

    def connect(self) -> None:
        raw_socket = socket.create_connection(
            (self._pinned_ip, self.port), self.timeout, self.source_address
        )
        # The validated IP is used for TCP; the original host is used for SNI
        # and certificate hostname verification by the default SSL context.
        try:
            self.sock = self._context.wrap_socket(raw_socket, server_hostname=self.host)
        except Exception:
            raw_socket.close()
            raise


def _download_url(url: str) -> tuple[str, bytes, str]:
    current_url = url.strip()
    for _ in range(6):
        host, port, ip = _check_public_http_url(current_url)
        parsed = urlparse(current_url)
        connection_class = _PinnedHTTPSConnection if parsed.scheme == "https" else _PinnedHTTPConnection
        connection = connection_class(host, port, ip)
        host_header = f"[{host}]" if ":" in host else host
        if port != (443 if parsed.scheme == "https" else 80):
            host_header += f":{port}"
        target = parsed.path or "/"
        if parsed.query:
            target += "?" + parsed.query
        try:
            connection.request("GET", target, headers={
                "Host": host_header,
                "User-Agent": "NotebookLM-App/1.0",
                "Accept": "text/html, text/plain, application/pdf",
                "Accept-Encoding": "identity",
            })
            response = connection.getresponse()
            if response.status in {301, 302, 303, 307, 308}:
                location = response.getheader("Location")
                if not location:
                    raise ValueError("URL redirect has no destination")
                current_url = urljoin(current_url, location)
                continue
            if not 200 <= response.status < 300:
                raise ValueError(f"Web server returned HTTP {response.status}")
            content_type = (response.getheader("Content-Type") or "text/html").split(";", 1)[0].lower()
            if content_type not in {"text/html", "application/xhtml+xml", "text/plain", "application/pdf"}:
                raise ValueError(f"Unsupported web content type: {content_type}")
            if (response.getheader("Content-Encoding") or "identity").lower() != "identity":
                raise ValueError("Web server returned unsupported compressed content")
            if response.length is not None and response.length > MAX_WEB_BYTES:
                raise ValueError("Web source exceeds the 5 MB download limit")
            deadline = time.monotonic() + 20
            parts = []
            size = 0
            while True:
                remaining = deadline - time.monotonic()
                if remaining <= 0:
                    raise ValueError("Web source download timed out")
                if connection.sock is not None:
                    connection.sock.settimeout(remaining)
                part = response.read(min(65536, MAX_WEB_BYTES - size + 1))
                if not part:
                    break
                size += len(part)
                if size > MAX_WEB_BYTES:
                    raise ValueError("Web source exceeds the 5 MB download limit")
                parts.append(part)
            return current_url, b"".join(parts), content_type
        except (OSError, http.client.HTTPException) as exc:
            raise ValueError(f"Could not download URL: {exc}") from exc
        finally:
            connection.close()
    raise ValueError("URL redirected too many times")


def ingest_url(store: NotebookStore, notebook_id: str, url: str,
               model_name: str = DEFAULT_MODEL) -> dict[str, Any]:
    final_url, body, content_type = _download_url(url)
    if content_type == "application/pdf":
        text = _extract_pdf(body)
        name = Path(urlparse(final_url).path).name or final_url
    elif content_type == "text/plain":
        text = _normalize_text(_decode_text(body))
        name = final_url
    else:
        try:
            from bs4 import BeautifulSoup
        except ImportError as exc:
            raise RuntimeError("Install beautifulsoup4 to ingest web pages") from exc
        # Pass decoded text so BeautifulSoup does not depend on whichever
        # optional charset detector happens to be installed in the runtime.
        soup = BeautifulSoup(_decode_text(body), "html.parser")
        name = (soup.title.get_text(" ", strip=True) if soup.title else "") or final_url
        for element in soup(["script", "style", "noscript", "nav", "footer", "header", "title"]):
            element.decompose()
        text = _normalize_text(soup.get_text("\n", strip=True))
    return ingest_text(store, notebook_id, name, text, "url", uri=final_url, model_name=model_name)
